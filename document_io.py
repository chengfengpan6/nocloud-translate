from pathlib import Path
from typing import List, Sequence, Tuple

from docx import Document


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx"}


def validate_file_extension(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS


def read_document(path: str | Path) -> Tuple[List[str], str]:
    file_path = Path(path)
    ext = file_path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("unsupported_format")

    if ext in {".txt", ".md"}:
        text = file_path.read_text(encoding="utf-8")
        blocks = text.splitlines()
        if text.endswith("\n"):
            blocks.append("")
        return blocks, ext

    document = Document(str(file_path))
    return [paragraph.text for paragraph in document.paragraphs], ext


def has_meaningful_text(blocks: Sequence[str]) -> bool:
    return any(block and block.strip() for block in blocks)


def build_output_path(original_path: str | Path, target_lang: str, output_ext: str | None = None) -> Path:
    original = Path(original_path)
    ext = output_ext or original.suffix.lower()
    if ext in {".txt", ".md"}:
        ext = ".txt"
    output_name = f"{original.stem}_translated_{target_lang}{ext}"
    return original.with_name(output_name)


def write_translated_document(
    original_path: str | Path,
    translated_blocks: Sequence[str],
    target_lang: str,
) -> Path:
    original = Path(original_path)
    ext = original.suffix.lower()
    output_path = build_output_path(original, target_lang, ext)

    if ext in {".txt", ".md"}:
        output_path.write_text("\n".join(translated_blocks), encoding="utf-8")
        return output_path

    if ext == ".docx":
        document = Document()
        for block in translated_blocks:
            document.add_paragraph(block)
        document.save(str(output_path))
        return output_path

    raise ValueError("unsupported_format")
